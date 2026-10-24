#!/usr/bin/env python3
"""Build a polished, editable DOCX from the same authoritative Group 2 Markdown."""
from __future__ import annotations
import re
import subprocess
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE

HERE=Path(__file__).resolve().parent
SOURCE=HERE/"Group_2_Final_Project_Report.md"
DEST=HERE/"Group_2_Final_Project_Report.docx"
NAVY=RGBColor(19,36,60);TEAL=RGBColor(12,98,115)
MUTE=RGBColor(86,104,117);WHITE=RGBColor(255,255,255)
LIGHT="EAF3F6";PALE="F5F8FA";DARK="12324B"

def hexcolor(rgb):
    return "".join(f"{int(c):02X}" for c in (rgb[0],rgb[1],rgb[2]))

def shading(cell, color):
    tcPr=cell._tc.get_or_add_tcPr()
    x=tcPr.find(qn("w:shd"))
    if x is None:x=OxmlElement("w:shd");tcPr.append(x)
    x.set(qn("w:fill"),color)

def cell_border(cell,bottom="D9E5EB"):
    pr=cell._tc.get_or_add_tcPr()
    borders=pr.find(qn("w:tcBorders"))
    if borders is None:borders=OxmlElement("w:tcBorders");pr.append(borders)
    for side in ("bottom",):
        b=borders.find(qn("w:"+side))
        if b is None:b=OxmlElement("w:"+side);borders.append(b)
        for k,v in {"val":"single","sz":"4","color":bottom}.items():b.set(qn("w:"+k),v)

def cell_margin(cell,top=95,left=100,bottom=90,right=100):
    tc=cell._tc;tcPr=tc.get_or_add_tcPr()
    mar=tcPr.find(qn("w:tcMar"))
    if mar is None:mar=OxmlElement("w:tcMar");tcPr.append(mar)
    for side,value in (("top",top),("start",left),("bottom",bottom),("end",right)):
        child=mar.find(qn("w:"+side))
        if child is None:child=OxmlElement("w:"+side);mar.append(child)
        child.set(qn("w:w"),str(value));child.set(qn("w:type"),"dxa")

def cant_split(row):
    pr=row._tr.get_or_add_trPr()
    if pr.find(qn("w:cantSplit")) is None:pr.append(OxmlElement("w:cantSplit"))

def repeat_header(row):
    pr=row._tr.get_or_add_trPr()
    if pr.find(qn("w:tblHeader")) is None:pr.append(OxmlElement("w:tblHeader"))

def field(par, instruction):
    run=par.add_run()
    a=OxmlElement("w:fldChar");a.set(qn("w:fldCharType"),"begin");run._r.append(a)
    r=par.add_run();b=OxmlElement("w:instrText");b.set(qn("xml:space"),"preserve")
    b.text=" "+instruction+" ";r._r.append(b)
    c=par.add_run();el=OxmlElement("w:fldChar");el.set(qn("w:fldCharType"),"end");c._r.append(el)

def text_style(style,name,size,color,bold=False):
    style.font.name=name;style.font.size=Pt(size);style.font.color.rgb=color;style.font.bold=bold

def after(element, other_element):
    element.addnext(other_element)

def build():
    subprocess.run([
        "pandoc",SOURCE.name,"--from=markdown+raw_tex+tex_math_dollars",
        "--to=docx","--standalone","--number-sections","--metadata=toc:false",
        "--resource-path=.:../reports","--output",DEST.name
    ],cwd=HERE,check=True)
    doc=Document(DEST)
    sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.left_margin=Inches(.85);sec.right_margin=Inches(.85)
    sec.top_margin=Inches(.84);sec.bottom_margin=Inches(.77)
    sec.header_distance=Inches(.31);sec.footer_distance=Inches(.33)
    sec.different_first_page_header_footer=True

    st=doc.styles
    for name in ["Normal","Body Text","First Paragraph","Block Text","Definition","Compact"]:
        if name in st:
            s=st[name];text_style(s,"Liberation Sans",9.5,NAVY)
            s.paragraph_format.space_before=Pt(0);s.paragraph_format.space_after=Pt(6.3)
            s.paragraph_format.line_spacing=1.17
            s.paragraph_format.widow_control=True
    for name,size,before,afterspace in [("Heading 1",15.0,16,6.5),("Heading 2",11.7,12,4.5),("Heading 3",10.0,8,3)]:
        if name in st:
            s=st[name];text_style(s,"Liberation Sans",size,TEAL,True)
            s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(afterspace)
            s.paragraph_format.keep_with_next=True;s.paragraph_format.keep_together=True
    for name,size,color in [("Title",25,NAVY),("Subtitle",12.0,MUTE),("Author",10,NAVY),("Date",9.0,MUTE)]:
        if name in st:
            s=st[name];text_style(s,"Liberation Sans",size,color,name=="Title")
            s.paragraph_format.space_after=Pt(11)
            s.paragraph_format.keep_with_next=True
    if "Image Caption" in st:
        text_style(st["Image Caption"],"Liberation Sans",8.25,MUTE)
        st["Image Caption"].paragraph_format.space_before=Pt(3)
        st["Image Caption"].paragraph_format.space_after=Pt(12)
        st["Image Caption"].paragraph_format.keep_with_next=False
    if "Caption" in st:
        text_style(st["Caption"],"Liberation Sans",8.1,MUTE)
    if "TOC Entry - Human" not in st:
        s=st.add_style("TOC Entry - Human",WD_STYLE_TYPE.PARAGRAPH)
        text_style(s,"Liberation Sans",10.5,NAVY)
        s.paragraph_format.space_before=Pt(0);s.paragraph_format.space_after=Pt(10)
        s.paragraph_format.left_indent=Inches(.18)

    # Cover page: accurate, original authorship and easily scanned project scope.
    paras=doc.paragraphs
    title=next(p for p in paras if p.style.name=="Title")
    kicker=title.insert_paragraph_before("THE UNIVERSITY OF TEXAS AT DALLAS  /  BUAN 6341")
    kicker.alignment=WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_before=Pt(92)
    kicker.paragraph_format.space_after=Pt(34)
    for r in kicker.runs:
        r.font.name="Liberation Sans";r.font.size=Pt(9);r.font.bold=True;r.font.color.rgb=TEAL
    for name in ("Title","Subtitle","Author","Date"):
        p=next((x for x in doc.paragraphs if x.style.name==name),None)
        if p:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    dt=next(p for p in doc.paragraphs if p.style.name=="Date")
    dt.paragraph_format.space_after=Pt(32)

    metric=doc.add_table(rows=2,cols=4)
    metric.autofit=True
    header=["DATA COVERAGE","MODEL SELECTION","EVALUATION DESIGN","HOLDOUT ACCURACY"]
    vals=["25,283 programs","4 model families","5 school-group folds","$6,957 MAE"]
    for i in range(4):
        for k,val in [(0,header[i]),(1,vals[i])]:
            c=metric.cell(k,i);c.text=val;c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shading(c,LIGHT if k==0 else PALE)
            cell_margin(c,130,80,110,80)
            for pp in c.paragraphs:
                pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in pp.runs:
                    run.font.name="Liberation Sans";run.font.size=Pt(7.4 if k==0 else 11.3)
                    run.bold=True;run.font.color.rgb=TEAL if k==0 else NAVY
    for rw in metric.rows:cant_split(rw)
    after(dt._p,metric._tbl)
    date_text=dt.text
    if "October" not in date_text:raise AssertionError("Metadata date missing")
    cover_note=doc.add_paragraph("RESEARCH REPORT  |  EDA, MODEL COMPARISON, FINE-TUNING, EXPLAINABILITY AND BUSINESS APPLICATION")
    cover_note.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cover_note.paragraph_format.space_before=Pt(20)
    for rr in cover_note.runs:rr.font.name="Liberation Sans";rr.font.size=Pt(8);rr.font.bold=True;rr.font.color.rgb=MUTE
    after(metric._tbl,cover_note._p)

    # Real contents, not an empty Word TOC field.
    toc=doc.add_paragraph("Contents",style="Heading 1")
    after(cover_note._p,toc._p)
    toc.paragraph_format.page_break_before=True
    toc.alignment=WD_ALIGN_PARAGRAPH.LEFT
    toc.paragraph_format.space_after=Pt(28)
    for rr in toc.runs:rr.font.name="Liberation Sans";rr.font.size=Pt(21);rr.font.bold=True;rr.font.color.rgb=NAVY
    headings=[p.text for p in doc.paragraphs if p.style.name=="Heading 1" and p._p is not toc._p]
    last=toc._p
    for s in headings:
        entry=doc.add_paragraph(s.strip(),style="TOC Entry - Human")
        after(last,entry._p);last=entry._p
    project_focus=next((p for p in doc.paragraphs if "Project focus." in p.text and p.style.name=="Block Text"),None)
    if project_focus:project_focus.paragraph_format.page_break_before=True
    else:raise AssertionError("Project focus not found")

    # Tables: compact academic style, repeated headers, sensible white space.
    for table in doc.tables:
        if table is metric:continue
        table.autofit=True
        for row_i,row in enumerate(table.rows):
            cant_split(row)
            if row_i==0:repeat_header(row)
            for cell in row.cells:
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if row_i==0:shading(cell,DARK)
                elif row_i%2==0:shading(cell,PALE)
                cell_border(cell);cell_margin(cell,95,110,90,110)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                    p.paragraph_format.line_spacing=1.09
                    for run in p.runs:
                        run.font.name="Liberation Sans";run.font.size=Pt(7.6 if row_i==0 else 8.2)
                        if row_i==0:run.font.bold=True;run.font.color.rgb=WHITE
                        else:run.font.color.rgb=NAVY

    # Figures: explicit consistent numbering for the editable artifact.
    figure_no=0
    for p in doc.paragraphs:
        if p.style.name=="Image Caption":
            figure_no+=1
            if not p.text.startswith("Figure "):
                p.runs[0].text="Figure "+str(figure_no)+". "+p.runs[0].text if p.runs else "Figure "+str(figure_no)+". "+p.text
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if figure_no!=16:raise AssertionError(f"Expected 16 embedded captioned figures, found {figure_no}")

    # Running metadata and actual numeric pages after the first cover page.
    hd=sec.header.paragraphs[0];hd.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    rr=hd.add_run("UT DALLAS  |  APPLIED MACHINE LEARNING  |  GROUP 2")
    rr.font.name="Liberation Sans";rr.font.size=Pt(7.6);rr.font.bold=True;rr.font.color.rgb=MUTE
    ft=sec.footer.paragraphs[0];ft.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=ft.add_run("COLLEGE MAJORS 2026  •  ")
    run.font.name="Liberation Sans";run.font.size=Pt(7.6);run.font.color.rgb=MUTE
    field(ft,"PAGE")

    for p in doc.paragraphs:
        if p.style.name in ("Heading 1","Heading 2","Heading 3"):
            p.paragraph_format.keep_with_next=True
        if p.style.name=="Heading 1" and p.text.strip().endswith("Conclusion"):
            p.paragraph_format.page_break_before=True
    doc.core_properties.title="College Majors 2026: Group 2 Final Project Report"
    doc.core_properties.subject="BUAN 6341 Group 2; exploratory analysis, model selection, tuning and responsible business case"
    doc.core_properties.author="UT Dallas BUAN 6341 - Group 2"
    doc.core_properties.keywords="College Majors 2026, Group 2, applied machine learning, model selection"
    doc.save(DEST)
    d=Document(DEST)
    print("DOCX generated:",DEST,"bytes",DEST.stat().st_size,
          "figures",len(d.inline_shapes),"tables",len(d.tables),
          "sections",len(d.sections),"headings",len(headings))
if __name__=="__main__":build()
